from __future__ import annotations

import json
import tempfile
import unittest
import sys
from pathlib import Path
from unittest import mock

import pymupdf
from docx import Document
from PIL import Image, ImageDraw

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from pdflib import detect, env, export, extract as extract_mod, layout
from pdflib.extract import _native_page
from pdflib import tables


def make_text_pdf(path: Path, pages=1):
    document = pymupdf.open()
    for index in range(pages):
        page = document.new_page(width=595, height=842)
        page.insert_text((50, 60), f"Document page {index + 1}", fontsize=20)
        page.insert_textbox(
            pymupdf.Rect(50, 90, 545, 220),
            "This is a native text document used to verify routing and reading order. "
            "The text layer is complete and should not trigger OCR.",
            fontsize=11,
        )
    document.save(path)
    document.close()


def make_scan_pdf(path: Path):
    source = pymupdf.open()
    page = source.new_page(width=595, height=842)
    page.insert_text((50, 80), "Scanned page test", fontsize=22)
    page.insert_textbox(
        pymupdf.Rect(50, 120, 545, 240),
        "This page is rendered to an image and inserted into a new PDF.",
        fontsize=11,
    )
    pixmap = page.get_pixmap(dpi=150, alpha=False)
    target = pymupdf.open()
    image_page = target.new_page(width=595, height=842)
    image_page.insert_image(image_page.rect, stream=pixmap.tobytes("png"))
    target.save(path)
    target.close()
    source.close()


def make_table_pdf(path: Path):
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)
    xs = [50, 150, 290, 430, 545]
    ys = [80, 110, 140, 170, 200, 230]
    for x in xs:
        page.draw_line((x, ys[0]), (x, ys[-1]), width=0.8)
    for y in ys:
        page.draw_line((xs[0], y), (xs[-1], y), width=0.8)
    rows = [
        ("No.", "Name", "Quantity", "Amount"),
        ("1", "Paper", "3", "12.00"),
        ("2", "Printing", "5", "25.00"),
        ("3", "Travel", "1", "80.00"),
        ("Total", "", "", "117.00"),
    ]
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            page.insert_text((xs[column_index] + 6, ys[row_index] + 22),
                             value, fontsize=10)
    document.save(path)
    document.close()


class CorePipelineTests(unittest.TestCase):
    def test_environment_home_is_user_cache_and_respects_override(self):
        with mock.patch.dict("os.environ", {"PDFSKILL_HOME": ""}):
            default_home = env.skill_home()
        self.assertEqual(default_home.name, "pdf-extract-skill")
        self.assertNotIn(".dsh", default_home.parts)
        with tempfile.TemporaryDirectory() as directory:
            with mock.patch.dict("os.environ", {"PDFSKILL_HOME": directory}):
                self.assertEqual(env.skill_home(), Path(directory))

    def test_conversion_report_requires_verification_offer(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf_path = root / "sample.pdf"
            out_dir = root / "out"
            make_text_pdf(pdf_path)
            report = extract_mod.run(
                pdf_path, out_dir, formats=("md",), conversion=True)
            prompt = "转换已完成。是否需要我继续校验并给出转换准确度？"
            self.assertEqual(report["verification_prompt"], prompt)
            self.assertEqual(report["next_action"]["type"], "ask_verification")
            stored = json.loads(
                (out_dir / "run_report.json").read_text(encoding="utf-8"))
            self.assertEqual(stored["verification_prompt"], prompt)

    def test_detection_routes_text_and_scan_pages(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            text_pdf = root / "text.pdf"
            scan_pdf = root / "scan.pdf"
            make_text_pdf(text_pdf)
            make_scan_pdf(scan_pdf)
            text_info = detect.detect(text_pdf)
            scan_info = detect.detect(scan_pdf)
            self.assertEqual(text_info["pages"][0]["route"], "native")
            self.assertEqual(text_info["recommended_route"], "native-flow")
            self.assertEqual(scan_info["pages"][0]["route"], "scan")
            self.assertEqual(scan_info["recommended_route"], "scan-ocr")

    def test_reading_order_keeps_columns_together(self):
        blocks = [
            {"type": "text", "text": "Title", "bbox": (40, 20, 300, 40)},
            {"type": "text", "text": "Left one", "bbox": (40, 60, 260, 80)},
            {"type": "text", "text": "Right one", "bbox": (320, 60, 540, 80)},
            {"type": "text", "text": "Left two", "bbox": (40, 100, 260, 120)},
            {"type": "text", "text": "Right two", "bbox": (320, 100, 540, 120)},
        ]
        ordered = [block["text"] for block in layout.reading_order(blocks)]
        self.assertLess(ordered.index("Left one"), ordered.index("Left two"))
        self.assertLess(ordered.index("Left two"), ordered.index("Right one"))
        self.assertLess(ordered.index("Right one"), ordered.index("Right two"))

    def test_native_table_is_accepted_and_flow_docx_keeps_text(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf_path = root / "table.pdf"
            docx_path = root / "table.docx"
            make_table_pdf(pdf_path)
            document = pymupdf.open(pdf_path)
            page = _native_page(document, 0)
            document.close()
            self.assertEqual(len(page["tables"]), 1)
            self.assertEqual(page["tables"][0][0][:4],
                             ["No.", "Name", "Quantity", "Amount"])
            info = {
                "pdf": str(pdf_path),
                "page_count": 1,
                "document_kind": "document",
                "route": {"recommended_route": "native-flow"},
                "warnings": [],
            }
            export.write_docx(docx_path, "table", info, [dict(page, width_pt=595,
                                                               height_pt=842)],
                              pdf_path=pdf_path, mode="flow")
            output = Document(docx_path)
            combined = "\n".join(p.text for p in output.paragraphs)
            cells = "\n".join(cell.text for table in output.tables
                              for row in table.rows for cell in row.cells)
            self.assertIn("Paper", combined + cells)

    def test_merged_sequence_column_uses_top_row_and_fills_missing_value(self):
        image = Image.new("RGB", (600, 360), "white")
        draw = ImageDraw.Draw(image)
        xs = [10, 70, 180, 390, 590]
        ys = [10, 60, 130, 200, 270, 340]
        for x in xs:
            draw.line((x, ys[0], x, ys[-1]), fill="black", width=2)
        for y in (ys[0], ys[1], ys[3], ys[-1]):
            draw.line((xs[0], y, xs[2], y), fill="black", width=2)
        for y in ys:
            draw.line((xs[2], y, xs[-1], y), fill="black", width=2)

        lines = [
            {"text": "序号", "cx": 40, "cy": 35, "x0": 20, "y0": 22,
             "x1": 60, "y1": 48},
            {"text": "单位名称", "cx": 125, "cy": 35, "x0": 85, "y0": 22,
             "x1": 165, "y1": 48},
            {"text": "坐落位置", "cx": 285, "cy": 35, "x0": 220, "y0": 22,
             "x1": 350, "y1": 48},
            {"text": "甲单位", "cx": 125, "cy": 130, "x0": 90, "y0": 105,
             "x1": 165, "y1": 155},
            {"text": "2", "cx": 40, "cy": 270, "x0": 32, "y0": 250,
             "x1": 48, "y1": 290},
            {"text": "乙单位", "cx": 125, "cy": 270, "x0": 90, "y0": 245,
             "x1": 165, "y1": 295},
            {"text": "位置一", "cx": 285, "cy": 95, "x0": 220, "y0": 75,
             "x1": 350, "y1": 115},
            {"text": "位置二", "cx": 285, "cy": 165, "x0": 220, "y0": 145,
             "x1": 350, "y1": 185},
            {"text": "位置三", "cx": 285, "cy": 235, "x0": 220, "y0": 215,
             "x1": 350, "y1": 255},
            {"text": "位置四", "cx": 285, "cy": 305, "x0": 220, "y0": 285,
             "x1": 350, "y1": 325},
        ]
        result = tables.reconstruct(lines, image)
        self.assertEqual(result["mode"], "rules")
        rows = result["rows"]
        self.assertEqual(rows[1][0], "1")
        self.assertEqual(rows[3][0], "2")
        self.assertEqual(rows[1][1], "甲单位")
        self.assertEqual(rows[3][1], "乙单位")
        self.assertTrue(result["quality"]["sequence_repair"]["inferred"])

    def test_sparse_complex_table_prefers_layout_word(self):
        rows = [["单位名称"] + [""] * 11]
        rows.extend([["基本信息", "统一社会信用代码"] + [""] * 10 for _ in range(11)])
        assessment = tables.assess_candidate(
            rows,
            {
                "mode": "rules",
                "quality": {
                    "fill_ratio": 0.18,
                    "grid_width_ratio": 0.9,
                    "grid_height_ratio": 0.8,
                },
            },
            source_route="scan",
            mode="auto",
        )
        self.assertTrue(assessment["accepted"])
        self.assertTrue(assessment["layout_preferred"])

    def test_table_docx_keeps_text_outside_table(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "table.docx"
            info = {
                "pdf": "sample.pdf",
                "page_count": 1,
                "document_kind": "scanned",
                "route": {"recommended_route": "scan-ocr"},
                "warnings": [],
            }
            page = {
                "page": 1,
                "tables": [[["字段", "内容"], ["姓名", ""]]],
                "outside_blocks": [{
                    "text": "表单主标题",
                    "placement": "before",
                    "bbox": [0, 0, 100, 20],
                }],
                "text": "",
                "width_pt": 595,
                "height_pt": 842,
            }
            export.write_docx(
                path, "sample", info, [page], pdf_path=Path("sample.pdf"),
                mode="table")
            document = Document(path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("表单主标题", text)

    def test_table_docx_applies_merged_regions(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "merged.docx"
            info = {
                "pdf": "sample.pdf",
                "page_count": 1,
                "document_kind": "scanned",
                "route": {"recommended_route": "scan-ocr"},
                "warnings": [],
            }
            rows = [["合并标题", "", "", ""],
                    ["甲", "", "乙", ""],
                    ["", "", "", ""],
                    ["", "", "", ""]]
            page = {
                "page": 1,
                "tables": [rows],
                "merges": [[
                    {"top": 0, "bottom": 1, "left": 0, "right": 4},
                    {"top": 1, "bottom": 4, "left": 0, "right": 1},
                ]],
                "outside_blocks": [],
                "text": "",
                "quality": {
                    "merged_grid_used": True,
                    "column_ratios": [0.25, 0.25, 0.25, 0.25],
                    "row_ratios": [0.25, 0.25, 0.25, 0.25],
                    "table_aspect": 1.0,
                },
                "width_pt": 595,
                "height_pt": 842,
            }
            export.write_docx(
                path, "sample", info, [page], pdf_path=Path("sample.pdf"),
                mode="table")
            table = Document(path).tables[0]
            self.assertIs(table.cell(0, 0)._tc, table.cell(0, 3)._tc)

    def test_blank_page_is_not_sent_to_ocr(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "blank.pdf"
            document = pymupdf.open()
            document.new_page(width=595, height=842)
            document.save(path)
            document.close()
            info = detect.detect(path)
            self.assertEqual(info["pages"][0]["route"], "blank")
            self.assertEqual(info["recommended_route"], "skip")


if __name__ == "__main__":
    unittest.main()
